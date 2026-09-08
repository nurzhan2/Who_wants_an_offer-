"""Will a machine be able to read this resume?

Phase 2 established that ``pdfplumber`` turns a two-column resume into this:

    ЗАРПЛАТА покрытие тестами с 30 до 78 процентов.
    от 1 800 000 KZT Стек: Python 3, FastAPI, PostgreSQL, Redis, Kafka, Docker
    НАВЫКИ Студия «Пиксель-Мираж» (фриланс, part-time) — удалённо

A sidebar heading welded to a sentence from the other column. This project
sidesteps that by handing the PDF to a model, which sees the layout. **An
employer's applicant tracking system cannot.** It reads exactly this text layer,
and the candidate never finds out why nobody called.

So the same machinery answers a question worth money to the person applying:
what does the robot see? Every check here works on the extracted text layer and
its coordinates — no LLM. Using a capable model to judge this would be modelling
the wrong reader; the point is to be as dumb as the parser on the other side.

The findings are deliberately specific. "Improve readability" is not actionable;
"your email is a picture, type it next to the icon" is.
"""

import io
import re
import statistics
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from itertools import pairwise
from typing import Any

import docx
import pdfplumber

from app.core.config import settings
from app.core.logging import get_logger
from app.schemas.ats import (
    ATSCoverage,
    ATSKeywords,
    ATSReport,
    DocumentKind,
    DocumentOrigin,
    Finding,
    FindingCode,
    Recoverable,
    Severity,
)
from app.schemas.llm import ProfileExtraction

logger = get_logger(__name__)

#: What each defect costs. The score is the arithmetic of the findings, never a
#: separate judgement, so it can always be explained by pointing at the list.
#:
#: Keyed by code rather than by severity, because two critical findings are not
#: equally bad. A flat severity penalty scored a file with no text layer at
#: 65/100 — a number that reads like a pass on a resume from which an employer's
#: parser extracts literally nothing. Severity says how to display a finding;
#: this says how much damage it does.
PENALTY: dict[FindingCode, int] = {
    # The parser gets nothing at all. There is no partial credit to give.
    FindingCode.NO_TEXT_LAYER: 100,
    # Text comes out, but shuffled: the parser reads a sentence assembled from
    # two columns and files the candidate under whatever it made of it.
    FindingCode.BROKEN_GLYPHS: 70,
    FindingCode.COLUMN_INTERLEAVING: 50,
    FindingCode.UNMAPPED_FONT: 40,
    # Everything parses except the one thing needed to reply.
    FindingCode.CONTACTS_NOT_TEXT: 35,
    # The parser reads the words but attributes them to the wrong job.
    FindingCode.CONTENT_LOST: 30,
    FindingCode.DATES_NOT_EXTRACTABLE: 25,
    FindingCode.TEXT_IN_TABLES: 15,
    FindingCode.MISSING_SECTIONS: 12,
    # Not a readability defect: this is the document trying to cheat the reader
    # it is being audited for. It costs more than a broken font because a
    # tracking system that catches it drops the candidate rather than the file,
    # and because we will not be the ones who taught them to do it.
    FindingCode.HIDDEN_TEXT: 60,
    FindingCode.KEYWORD_STUFFING: 30,
    # The requirements this variant does not name. Small: the document is
    # readable and true, it is just answering less of the posting than it could.
    FindingCode.REQUIREMENTS_NOT_NAMED: 10,
    FindingCode.DATE_FORMAT_MIXED: 8,
    FindingCode.LENGTH_OUT_OF_RANGE: 5,
    # Costs nothing. A break in someone's employment is a fact about their life,
    # not a defect in their file, and the only thing the audit has to say about
    # it is what the parser will compute. See :func:`check_gaps`.
    FindingCode.EMPLOYMENT_GAP: 0,
    FindingCode.FORMAT_NOT_PDF: 0,
}

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]{2,}")
#: Loose about formatting, strict about length. An ATS regex is loose too, and
#: the question is only whether *something* phone-shaped survives into the text
#: layer. The digit floor and the missing dash variants are what stop a date
#: range from answering that question: "04.2022 — 09.2023" is a run of digits,
#: spaces and a dash, and a permissive pattern reads it as a phone number — so
#: a resume whose real phone is an image would quietly pass this check.
PHONE = re.compile(r"\+?\d[\d\s()-]{7,}\d")
#: Shortest national number worth calling, minus nothing. Kazakh and Russian
#: numbers are 11 digits, EU ones 9 and up.
MIN_PHONE_DIGITS = 9
#: pdfminer emits this when a font carries no ToUnicode map: the glyph is drawn
#: but there is no character behind it.
CID_GLYPH = re.compile(r"\(cid:\d+\)")

#: Ways a resume writes an employment period. An ATS keys seniority off these,
#: so a resume whose dates survive only as an image reads as no experience at
#: all. Loose about separators, strict about the year: a bare "2019" inside a
#: sentence is not a period.
DATE_TOKEN = re.compile(
    r"(?:\b(?:0?[1-9]|1[0-2])[./-](?:19|20)\d{2}\b)"
    r"|(?:\b(?:19|20)\d{2}[./-](?:0?[1-9]|1[0-2])\b)"
    r"|(?:\b(?:янв|фев|мар|апр|мая|май|июн|июл|авг|сен|окт|ноя|дек"
    r"|jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[а-яa-z]*\.?\s+(?:19|20)\d{2}\b)",
    re.IGNORECASE,
)
#: Below this a resume has no employment history a parser can date.
MIN_DATE_TOKENS = 2
#: Anything above this in a two-number date is the year, not the month.
MONTHS_IN_YEAR = 12

#: Section headings a parser looks for, in the two languages this market writes
#: resumes in. Missing them does not break extraction, it breaks *sectioning* —
#: the ATS cannot tell which dates belong to jobs and which to degrees.
SECTION_PATTERNS: dict[str, re.Pattern[str]] = {
    "experience": re.compile(
        r"\b(experience|employment|work\s+history)\b|опыт\s+работы|опыт", re.IGNORECASE
    ),
    "education": re.compile(r"\b(education|degree)\b|образование", re.IGNORECASE),
    "skills": re.compile(r"\b(skills|technical\s+skills|stack)\b|навыки|стек", re.IGNORECASE),
}

#: Characters a Russian or English resume is made of. Anything far outside this
#: is a font that did not survive extraction.
EXPECTED_CHARS = re.compile(r"[a-zA-Zа-яА-ЯёЁ0-9\s\.,;:!?()\[\]{}«»\"'`@#№%&*+/\\|_—–\-…€$₸₽]")


#: A check reads the pages and either reports something or stays quiet.
type Check = Callable[[Sequence[PageFacts]], Finding | None]


@dataclass(frozen=True, slots=True)
class Column:
    """A vertical band of the page that words line up in."""

    left: float
    right: float


@dataclass(frozen=True, slots=True)
class PageFacts:
    """Everything the checks need from one page, gathered once."""

    words: list[dict[str, Any]]
    text: str
    width: float
    columns: list[Column]
    mixed_lines: list[str]
    line_count: int
    table_words: int
    chars: list[dict[str, Any]]
    #: Filled shapes on the page. Only :func:`check_hidden_text` reads them, and
    #: only to answer one question: white text on a dark banner is a design, not
    #: a hidden keyword block, and without the shape under it the two are the
    #: same character. Defaulted because a DOCX and a TXT have no shapes.
    rects: list[dict[str, Any]] = field(default_factory=list)


def detect_columns(words: Sequence[dict[str, Any]], width: float) -> list[Column]:
    """Vertical bands the words start in.

    Clustering on the left edge rather than the whole box: a sidebar and a body
    column differ by where their lines *begin*, and using the full extent would
    merge them as soon as one long line crossed the gutter.

    A band has to hold a real share of the page's words to count. Otherwise a
    page number in a corner reads as a second column.
    """
    if not words:
        return []
    starts = sorted(word["x0"] for word in words)
    gap = width * settings.ats_column_gap_ratio
    clusters: list[list[float]] = [[starts[0]]]
    for value in starts[1:]:
        if value - clusters[-1][-1] > gap:
            clusters.append([value])
        else:
            clusters[-1].append(value)

    floor = len(words) * settings.ats_min_column_share
    return [Column(left=min(c), right=max(c)) for c in clusters if len(c) >= floor]


def _column_of(word: dict[str, Any], columns: Sequence[Column]) -> int:
    """Which band a word starts in, or -1 when it starts between them."""
    for index, column in enumerate(columns):
        if column.left - 1 <= word["x0"] <= column.right + 1:
            return index
    return -1


def _visual_lines(words: Sequence[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    """Group words into the rows a reader sees.

    Bucketed on the rounded top coordinate: text extraction gives every word its
    own baseline, and words on one visual line differ by fractions of a point.
    """
    rows: dict[int, list[dict[str, Any]]] = {}
    for word in words:
        rows.setdefault(round(word["top"] / 3), []).append(word)
    return [sorted(row, key=lambda w: w["x0"]) for _, row in sorted(rows.items())]


def find_mixed_lines(
    words: Sequence[dict[str, Any]], columns: Sequence[Column]
) -> tuple[list[str], int]:
    """Lines that read across two columns, and how many lines there were.

    This is the interleaving, made visible. A line holding words from two bands
    is one the extractor will emit as a single run of text, gluing a sidebar
    heading to a sentence that belongs somewhere else entirely.
    """
    lines = _visual_lines(words)
    if len(columns) < 2:
        return [], len(lines)

    mixed: list[str] = []
    for line in lines:
        bands = {_column_of(word, columns) for word in line} - {-1}
        if len(bands) >= 2:
            mixed.append(" ".join(word["text"] for word in line))
    return mixed, len(lines)


def _count_table_words(page: pdfplumber.page.Page, words: Sequence[dict[str, Any]]) -> int:
    """Words that sit inside a detected table.

    A skills grid drawn as a table, or a "proficiency" bar chart with labels, is
    the classic way to lose an entire skill list: the words are there, and they
    come out in an order that means nothing.
    """
    try:
        tables = page.find_tables()
    except Exception:  # pdfplumber's table finder is best-effort on odd files
        return 0

    boxes = [table.bbox for table in tables]
    if not boxes:
        return 0
    return sum(
        1
        for word in words
        for (x0, top, x1, bottom) in boxes
        if x0 <= word["x0"] <= x1 and top <= word["top"] <= bottom
    )


def gather(content: bytes) -> list[PageFacts]:
    """Read every page once, so each check does not re-open the file."""
    pages: list[PageFacts] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            words = page.extract_words()
            columns = detect_columns(words, page.width)
            mixed, line_count = find_mixed_lines(words, columns)
            pages.append(
                PageFacts(
                    words=words,
                    text=page.extract_text() or "",
                    width=page.width,
                    columns=columns,
                    mixed_lines=mixed,
                    line_count=line_count,
                    table_words=_count_table_words(page, words),
                    chars=page.chars,
                    rects=page.rects,
                )
            )
    return pages


def _finding(
    code: FindingCode,
    severity: Severity,
    title: str,
    explanation: str,
    fix: str,
    example: str | None = None,
) -> Finding:
    """One finding, with its penalty read off :data:`PENALTY`."""
    return Finding(
        code=code,
        severity=severity,
        title=title,
        explanation=explanation,
        example_fragment=example,
        fix=fix,
        penalty=PENALTY[code],
    )


def check_text_layer(pages: Sequence[PageFacts]) -> Finding | None:
    """Is there any text at all, or is this a picture of a resume?"""
    total = sum(len(page.text.strip()) for page in pages)
    if total >= settings.ats_min_text_chars:
        return None
    return _finding(
        FindingCode.NO_TEXT_LAYER,
        Severity.CRITICAL,
        "В файле нет текстового слоя",
        f"Из файла извлекается {total} символов. Это скан или картинка: "
        "робот работодателя не прочитает оттуда ни одного слова и увидит пустое резюме.",
        "Экспортируй резюме в PDF из текстового редактора, а не сканируй "
        "распечатку и не сохраняй как изображение.",
    )


def check_columns(pages: Sequence[PageFacts]) -> Finding | None:
    """Do the columns interleave when read as text?"""
    mixed = [line for page in pages for line in page.mixed_lines]
    lines = sum(page.line_count for page in pages)
    if not lines or not mixed:
        return None

    share = len(mixed) / lines
    if share < settings.ats_mixed_line_ratio:
        return None

    example = "\n".join(line[:100] for line in mixed[:3])
    return _finding(
        FindingCode.COLUMN_INTERLEAVING,
        Severity.CRITICAL,
        "Колонки перемешиваются при чтении",
        f"Резюме свёрстано в несколько колонок, и {share:.0%} строк текстового слоя "
        "склеивают боковую колонку с основной. Робот читает страницу построчно "
        "поперёк колонок и получает именно это:",
        "Перевёрстай в одну колонку. Это единственная надёжная починка: "
        "боковая панель с навыками и контактами всегда будет склеиваться "
        "с текстом рядом.",
        example,
    )


def check_contacts(pages: Sequence[PageFacts]) -> Finding | None:
    """Can a machine find an email and a phone number?"""
    text = "\n".join(page.text for page in pages)
    has_email = bool(EMAIL.search(text))
    has_phone = any(
        sum(char.isdigit() for char in match) >= MIN_PHONE_DIGITS for match in PHONE.findall(text)
    )
    if has_email and has_phone:
        return None

    missing = [
        name for name, present in (("почту", has_email), ("телефон", has_phone)) if not present
    ]
    return _finding(
        FindingCode.CONTACTS_NOT_TEXT,
        Severity.CRITICAL,
        f"Робот не находит {' и '.join(missing)}",
        "В текстовом слое нет ничего похожего на "
        f"{' и '.join(missing)}. Обычно это значит, что контакты нарисованы "
        "иконками или вставлены картинкой. Резюме может пройти отбор, "
        "и связаться с тобой будет нечем.",
        "Впиши почту и телефон обычным текстом — рядом с иконкой или вместо неё. "
        "Иконка остаётся, текст добавляется.",
    )


def check_glyphs(pages: Sequence[PageFacts]) -> Finding | None:
    """Does the text come out as characters, or as noise?"""
    text = "".join(page.text for page in pages)
    stripped = "".join(text.split())
    if not stripped:
        return None

    recognised = len(EXPECTED_CHARS.findall(stripped))
    broken = 1 - recognised / len(stripped)
    if broken < settings.ats_broken_glyph_ratio:
        return None

    sample = next(
        (line for line in text.splitlines() if line.strip() and not EXPECTED_CHARS.match(line)),
        None,
    )
    return _finding(
        FindingCode.BROKEN_GLYPHS,
        Severity.CRITICAL,
        "Текст извлекается искажённым",
        f"{broken:.0%} символов текстового слоя не опознаются как буквы или цифры. "
        "Шрифт нарисован правильно, но за глифами нет символов — робот получает мусор.",
        "Пересохрани PDF с внедрением шрифтов, либо смени шрифт на "
        "стандартный. Проверить просто: выдели текст в резюме и вставь "
        "в блокнот — увидишь то же, что увидит робот.",
        (sample or "")[:100] or None,
    )


def check_unmapped_fonts(pages: Sequence[PageFacts]) -> Finding | None:
    """Are there glyphs with no character behind them?"""
    text = "".join(page.text for page in pages)
    hits = CID_GLYPH.findall(text)
    if not hits:
        return None
    return _finding(
        FindingCode.UNMAPPED_FONT,
        Severity.CRITICAL,
        "Шрифт без таблицы соответствия символов",
        f"В извлечённом тексте {len(hits)} мест вида (cid:NN). Шрифт вставлен без "
        "таблицы ToUnicode: буквы видны человеку и отсутствуют для машины.",
        "Пересохрани PDF из исходника с внедрением шрифтов, или замени шрифт на системный.",
        " ".join(hits[:8]),
    )


def check_tables(pages: Sequence[PageFacts]) -> Finding | None:
    """How much of the content lives inside a table?"""
    inside = sum(page.table_words for page in pages)
    total = sum(len(page.words) for page in pages)
    if not total:
        return None

    share = inside / total
    if share < settings.ats_table_word_ratio:
        return None
    return _finding(
        FindingCode.TEXT_IN_TABLES,
        Severity.WARNING,
        "Содержимое свёрстано таблицами",
        f"{share:.0%} слов находятся внутри таблиц. Многие ATS читают таблицу "
        "по ячейкам без учёта смысла, и список навыков или дат превращается "
        "в перемешанный набор слов.",
        "Замени табличную вёрстку обычными абзацами и списками. "
        "Особенно это касается блока навыков и шкал «уровень владения».",
    )


def detect_sections(pages: Sequence[PageFacts]) -> list[str]:
    """Which standard sections a parser would find, by the name it looks for."""
    text = "\n".join(page.text for page in pages)
    return sorted(name for name, pattern in SECTION_PATTERNS.items() if pattern.search(text))


def check_dates(pages: Sequence[PageFacts]) -> Finding | None:
    """Can a machine find employment periods, or only see them as a picture?"""
    text = "\n".join(page.text for page in pages)
    found = DATE_TOKEN.findall(text)
    if len(found) >= MIN_DATE_TOKENS:
        return None
    return _finding(
        FindingCode.DATES_NOT_EXTRACTABLE,
        Severity.WARNING,
        "Робот не находит периоды работы",
        f"В текстовом слое найдено дат в узнаваемом формате: {len(found)}. "
        "ATS считает стаж по датам рядом с местами работы; без них резюме "
        "выглядит как опыт длиной ноль лет, каким бы он ни был на самом деле.",
        "Пиши периоды текстом в формате ММ.ГГГГ — «04.2022 — 09.2023» или "
        "«04.2022 — по настоящее время». Не выноси даты в графику и не "
        "оставляй только годы без месяцев.",
    )


def check_sections(pages: Sequence[PageFacts]) -> Finding | None:
    """Are the standard section headings there to be found?"""
    text = "\n".join(page.text for page in pages)
    missing = sorted(name for name, pattern in SECTION_PATTERNS.items() if not pattern.search(text))
    if not missing:
        return None

    russian = {"experience": "опыт работы", "education": "образование", "skills": "навыки"}
    return _finding(
        FindingCode.MISSING_SECTIONS,
        Severity.WARNING,
        "Нет узнаваемых заголовков разделов",
        "Не найдены заголовки: " + ", ".join(russian[name] for name in missing) + ". "
        "Робот разбирает резюме по секциям; без заголовков он не отличит "
        "даты работы от дат учёбы.",
        "Добавь простые заголовки «Опыт работы», «Образование», «Навыки» — "
        "именно словами, а не только визуальным выделением.",
    )


# ── the checks that read the document as a document ────────────────────
#
# Everything above asks whether the text comes out. These ask whether what
# comes out is shaped like a resume a parser can file, and — the one check here
# that is not about readability at all — whether the document is trying to
# cheat the reader it is being audited for.

#: Codepoints that occupy no space and draw nothing. In a resume's text layer
#: they arrive one of two ways: pasted in from a web page, or sprinkled through
#: a keyword block to stop a human reader noticing it. Either way an employer's
#: parser reads them as part of the words around them.
INVISIBLE_CHARS = re.compile(r"[\u00ad\u200b-\u200f\u202a-\u202e\u2060-\u2064\ufeff]")
#: Below this many, an invisible codepoint is a copy-paste artefact rather than
#: a technique. Three is the smallest count that cannot be one stray paste.
MIN_INVISIBLE = 3

#: Words too common to mean anything when repeated. Everything shorter than
#: three characters is already excluded, which covers most of the Russian ones.
STOPWORDS = frozenset(
    [
        "без",
        "были",
        "было",
        "была",
        "будет",
        "всё",
        "где",
        "для",
        "если",
        "есть",
        "его",
        "ещё",
        "или",
        "их",
        "как",
        "когда",
        "который",
        "которые",
        "меня",
        "над",
        "них",
        "ним",
        "она",
        "они",
        "при",
        "про",
        "свою",
        "так",
        "там",
        "того",
        "тоже",
        "только",
        "что",
        "чтобы",
        "это",
        "этом",
        "and",
        "are",
        "but",
        "для",
        "for",
        "from",
        "has",
        "have",
        "his",
        "her",
        "its",
        "not",
        "our",
        "that",
        "the",
        "them",
        "then",
        "they",
        "this",
        "was",
        "were",
        "what",
        "when",
        "which",
        "with",
        "you",
        "your",
        "года",
        "год",
        "лет",
        "компания",
        "компании",
        "опыт",
        "работа",
        "работы",
        "проект",
        "проекта",
    ]
)

#: A token repeated more often than this in one document is being repeated on
#: purpose. Measured against the fixtures: the busiest legitimate repeat in a
#: real one-page resume is a technology named once per job, which is four.
KEYWORD_REPEAT_LIMIT = 12
#: The same token this many times in a row is a keyword block whatever its
#: total count is — "python python python python" is not a sentence.
CONSECUTIVE_REPEAT_LIMIT = 4

#: How a document writes a date, by the shape of it. Mixing two of these is not
#: a style problem: a parser recognises the formats it was written for, and a
#: resume that uses two has a chance of being read under only one of them.
DATE_STYLES: dict[str, re.Pattern[str]] = {
    "MM.YYYY": re.compile(r"\b(?:0?[1-9]|1[0-2])[./-](?:19|20)\d{2}\b"),
    "YYYY-MM": re.compile(r"\b(?:19|20)\d{2}[./-](?:0?[1-9]|1[0-2])\b"),
    "месяц ГГГГ": re.compile(
        r"\b(?:янв|фев|мар|апр|мая|май|июн|июл|авг|сен|окт|ноя|дек"
        r"|jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[а-яa-z]*\.?\s+(?:19|20)\d{2}\b",
        re.IGNORECASE,
    ),
}
#: One date in a second format is a typo; two is a habit the document is built
#: on, and only the second is worth telling somebody to go and fix.
MIN_MINORITY_DATES = 2

#: Months of silence a parser reads as a break in employment.
GAP_MONTHS = 12


def _luminance(color: object) -> float | None:
    """How light a pdfplumber colour is, 0.0 black to 1.0 white.

    Handles the three shapes a PDF states a colour in — grey, RGB, CMYK — and
    returns ``None`` for anything else, including the ``None`` a page uses to
    mean "whatever was set before". Guessing at an unknown colour space would
    mean either inventing hidden text or missing it, and both are worse than
    saying the check could not read this character.
    """
    if isinstance(color, (int, float)) and not isinstance(color, bool):
        return max(0.0, min(1.0, float(color)))
    if not isinstance(color, (list, tuple)) or not color:
        return None
    try:
        values = [float(component) for component in color]
    except (TypeError, ValueError):
        return None
    if len(values) == 1:
        return max(0.0, min(1.0, values[0]))
    if len(values) == 3:
        red, green, blue = values
        return max(0.0, min(1.0, 0.299 * red + 0.587 * green + 0.114 * blue))
    if len(values) == 4:
        cyan, magenta, yellow, black = values
        red = (1 - cyan) * (1 - black)
        green = (1 - magenta) * (1 - black)
        blue = (1 - yellow) * (1 - black)
        return max(0.0, min(1.0, 0.299 * red + 0.587 * green + 0.114 * blue))
    return None


def _dark_boxes(page: PageFacts) -> list[tuple[float, float, float, float]]:
    """Filled shapes dark enough for white text to be readable on them.

    Without this the check would report every resume with a dark header band as
    hiding text, which is a false critical on a common and entirely honest
    design — and a false critical tells somebody to rebuild a file that was
    fine.
    """
    boxes: list[tuple[float, float, float, float]] = []
    for rect in page.rects:
        light = _luminance(rect.get("non_stroking_color"))
        if light is None or light > settings.ats_dark_fill_luminance:
            continue
        try:
            boxes.append(
                (float(rect["x0"]), float(rect["top"]), float(rect["x1"]), float(rect["bottom"]))
            )
        except (KeyError, TypeError, ValueError):  # a shape with no usable box
            continue
    return boxes


def _is_hidden(char: dict[str, Any], dark: Sequence[tuple[float, float, float, float]]) -> bool:
    """Is this character drawn so a person cannot see it but a parser can?

    Two ways: painted the colour of the page, or set so small it reads as a
    line of dust. Both are visible to text extraction and to nothing else.
    """
    if char.get("text", "").isspace():
        return False
    size = char.get("size")
    if isinstance(size, (int, float)) and 0 < float(size) < settings.ats_min_font_size:
        return True

    light = _luminance(char.get("non_stroking_color"))
    if light is None or light < settings.ats_invisible_luminance:
        return False
    try:
        x = (float(char["x0"]) + float(char["x1"])) / 2
        y = (float(char["top"]) + float(char["bottom"])) / 2
    except (KeyError, TypeError, ValueError):
        return True
    return not any(x0 <= x <= x1 and top <= y <= bottom for x0, top, x1, bottom in dark)


def check_hidden_text(pages: Sequence[PageFacts]) -> Finding | None:
    """Is anything in this document written to be read by machines only?

    This is the one check here that is not about whether a parser copes. White
    text on white paper, a two-point keyword block, zero-width characters
    threaded through a paragraph — none of them break extraction, all of them
    are aimed at it, and every applicant tracking system worth the name has been
    catching them for a decade. Being caught does not filter the file, it
    filters the person.

    So it is reported as a defect, in those words, and the fix says to delete
    it. The audit's whole claim on somebody's trust is that it tells them what
    the machine sees; an audit that noticed this and stayed quiet — or worse,
    framed it as working — would be teaching the trick it was asked to detect.
    """
    hidden: list[str] = []
    for page in pages:
        dark = _dark_boxes(page)
        hidden.extend(char.get("text", "") for char in page.chars if _is_hidden(char, dark))

    text = "".join(page.text for page in pages)
    invisible = INVISIBLE_CHARS.findall(text)
    if len(hidden) < settings.ats_hidden_char_limit and len(invisible) < MIN_INVISIBLE:
        return None

    if len(hidden) >= settings.ats_hidden_char_limit:
        what = (
            f"{len(hidden)} символов набраны цветом фона или размером, который человек не разглядит"
        )
        example = "".join(hidden)[:100]
    else:
        what = f"{len(invisible)} невидимых символов вставлены внутрь текста"
        example = None
    return _finding(
        FindingCode.HIDDEN_TEXT,
        Severity.CRITICAL,
        "В документе есть скрытый текст",
        f"{what}. Робот их читает, человек — нет. Это распознаётся как попытка "
        "обмануть отбор: системы отслеживания кандидатов сравнивают видимый "
        "слой с текстовым и помечают такие резюме, после чего отклоняют не "
        "файл, а кандидата.",
        "Удали скрытый блок целиком. Ключевые слова работают только тогда, "
        "когда они написаны в тексте, который человек прочитает и сможет "
        "подтвердить на собеседовании.",
        example,
    )


def _tokens(pages: Sequence[PageFacts]) -> list[str]:
    """Words of the document, folded, short ones and stopwords dropped."""
    text = "\n".join(page.text for page in pages)
    return [
        word
        for word in re.findall(r"[^\W_]{3,}", text.casefold(), re.UNICODE)
        if word not in STOPWORDS
    ]


def check_stuffing(pages: Sequence[PageFacts]) -> Finding | None:
    """Is the same word repeated past the point of meaning anything?

    The other half of the boundary this audit refuses to cross. Keyword stuffing
    is the advice the internet gives, it is what a generator optimising for a
    coverage number would converge on, and it is what this check exists to stop
    — including on documents this system wrote itself.
    """
    tokens = _tokens(pages)
    if not tokens:
        return None

    counts: dict[str, int] = {}
    run_token, run, longest_run, worst_run_token = "", 0, 0, ""
    for token in tokens:
        counts[token] = counts.get(token, 0) + 1
        run = run + 1 if token == run_token else 1
        run_token = token
        if run > longest_run:
            longest_run, worst_run_token = run, token

    repeated = sorted(
        (token for token, count in counts.items() if count > KEYWORD_REPEAT_LIMIT),
        key=lambda token: -counts[token],
    )
    if not repeated and longest_run < CONSECUTIVE_REPEAT_LIMIT:
        return None

    if repeated:
        detail = ", ".join(f"«{token}» — {counts[token]} раз" for token in repeated[:5])
    else:
        detail = f"«{worst_run_token}» — {longest_run} раз подряд"
    return _finding(
        FindingCode.KEYWORD_STUFFING,
        Severity.WARNING,
        "Ключевые слова повторяются набивкой",
        f"В тексте {detail}. Так выглядит не резюме, а список для робота: "
        "современные ATS считают частоту и понижают документы, где она "
        "неестественная, а человек, открывший файл следом, прочитает то же самое.",
        "Оставь каждое название столько раз, сколько его требует смысл — "
        "обычно по одному разу на место работы, где навык применялся. "
        "Повторение не увеличивает совпадение, а помечает документ.",
        detail[:200],
    )


def check_date_format(pages: Sequence[PageFacts]) -> Finding | None:
    """Does the document write its periods one way, or several?"""
    text = "\n".join(page.text for page in pages)
    counts = {style: len(pattern.findall(text)) for style, pattern in DATE_STYLES.items()}
    used = {style: count for style, count in counts.items() if count}
    if len(used) < 2:
        return None
    # The dominant style is the document's; the question is whether the others
    # are a habit or a typo. One stray date is not worth an instruction to go
    # and edit a file.
    minority = sorted(used.values())[:-1]
    if max(minority) < MIN_MINORITY_DATES:
        return None

    listing = ", ".join(f"{style} — {count}" for style, count in used.items())
    return _finding(
        FindingCode.DATE_FORMAT_MIXED,
        Severity.WARNING,
        "Даты записаны в разных форматах",
        f"В резюме встречаются форматы: {listing}. Парсер разбирает даты по "
        "шаблонам и обычно знает не все: то, что записано вторым форматом, "
        "рискует не превратиться в период работы, и стаж посчитается меньше.",
        "Приведи все периоды к одному виду — «04.2022 — 09.2023», "
        "«04.2022 — по настоящее время». Включая даты образования.",
        listing,
    )


def check_length(pages: Sequence[PageFacts]) -> Finding | None:
    """Is there enough here to match against, and not so much it gets cut?"""
    words = _word_count(pages)
    if not words:
        return None
    if words < settings.ats_min_resume_words:
        return _finding(
            FindingCode.LENGTH_OUT_OF_RANGE,
            Severity.WARNING,
            "Резюме слишком короткое для сопоставления",
            f"В документе {words} слов. Отбор по ключевым словам работает с тем, "
            "что написано: у короткого резюме почти нет совпадений — не потому, "
            "что опыта нет, а потому что он не назван.",
            "Опиши каждое место работы задачами и технологиями, а не одной "
            "строкой должности. Названия инструментов пиши словами.",
        )
    if words > settings.ats_max_resume_words:
        return _finding(
            FindingCode.LENGTH_OUT_OF_RANGE,
            Severity.WARNING,
            "Резюме длиннее, чем читает робот",
            f"В документе {words} слов. Часть систем обрезает текст при импорте, "
            "и обрезается всегда конец — то есть ранний опыт и образование.",
            "Сократи до двух страниц: подробно последние места работы, "
            "остальные — строкой с датами и должностью.",
        )
    return None


def check_gaps(extraction: ProfileExtraction) -> Finding | None:
    """Periods a parser will read as a break in employment.

    Reported at zero cost and at ``info``, and the wording is the reason this
    check is written separately from the rest. A gap is usually a fact about
    somebody's life — a child, an illness, a year of study, a country change —
    and it is not this audit's business to imply it should be papered over. The
    only thing being said is what the arithmetic on the other side will produce,
    and the only fix offered is to date work that is already in the document.
    """
    dated = sorted(
        (
            (start, year_month(str(period.end)) if period.end else None, period.company)
            for period in extraction.work_periods
            if period.start and (start := year_month(str(period.start))) is not None
        ),
        key=lambda item: item[0],
    )
    if len(dated) < 2:
        return None

    gaps: list[str] = []
    for (_, end, company), (start, _, later) in pairwise(dated):
        if end is None:
            continue
        months = (start[0] - end[0]) * MONTHS_IN_YEAR + (start[1] - end[1])
        if months > GAP_MONTHS:
            gaps.append(f"{company} → {later}: {months} мес.")
    if not gaps:
        return None

    return _finding(
        FindingCode.EMPLOYMENT_GAP,
        Severity.INFO,
        "Между периодами работы есть промежутки",
        "Робот считает стаж как сумму периодов, поэтому эти промежутки в стаж "
        "не войдут: " + "; ".join(gaps) + ". Это не дефект файла — так "
        "посчитает любая система, читающая даты.",
        "Если в эти месяцы была работа, учёба, фриланс или свой проект — "
        "добавь их с датами в том же формате. Если не было, ничего "
        "исправлять не нужно: промежуток в биографии не чинится резюме.",
    )


def check_requirements(keywords: ATSKeywords) -> Finding | None:
    """Requirements the candidate holds and this variant does not name.

    Only the middle bucket produces a finding. The requirements nobody holds are
    reported in the keyword list and generate no advice at all, because the only
    advice available would be to claim them.
    """
    unstated = keywords.unstated
    if not unstated:
        return None

    named = ", ".join(
        f"«{item.requirement}»" + (f" (в резюме: «{item.found_as}»)" if item.found_as else "")
        for item in unstated[:8]
    )
    return _finding(
        FindingCode.REQUIREMENTS_NOT_NAMED,
        Severity.WARNING,
        f"Не названы дословно: {len(unstated)} требований из списка вакансии",
        "Эти навыки есть в профиле, но в этом варианте документа они не "
        f"написаны так, как их ищет работодатель: {named}. Фильтр ищет точные "
        "строки, поэтому «постгрес» и «PostgreSQL» для него разные вещи.",
        "Перегенерируй вариант под эту вакансию: он назовёт эти навыки "
        "формулировками вакансии там, где они действительно были в работе. "
        "Ничего нового при этом не появляется — всё перечисленное уже в профиле.",
    )


#: Checks that need a PDF's geometry, in the order they are reported.
PDF_CHECKS = (
    check_text_layer,
    check_columns,
    check_unmapped_fonts,
    check_glyphs,
    check_hidden_text,
    check_contacts,
    check_dates,
    check_date_format,
    check_tables,
    check_sections,
    check_stuffing,
    check_length,
)
#: Checks that work on any text, whatever produced it.
#: What can be asked of text with no page behind it.
TEXT_CHECKS = (
    check_contacts,
    check_glyphs,
    check_hidden_text,
    check_dates,
    check_date_format,
    check_sections,
    check_stuffing,
    check_length,
)

#: A DOCX adds the one layout question worth asking of it. Kept separate from
#: TEXT_CHECKS so a plain .txt does not claim to have been checked for tables
#: it cannot have.
DOCX_CHECKS = (*TEXT_CHECKS, check_tables)

#: What may be asked of a cover letter.
#:
#: Short, and the shortness is the point. A letter has no employment section, no
#: date column and no second page, so running the resume checks over one would
#: report four defects on a perfectly good letter — an audit that cries wolf
#: about the wrong document is one nobody reads on the right one. What survives
#: is the pair that is about conduct rather than layout, and those two matter
#: more here than anywhere else: this is the document *we* wrote.
LETTER_CHECKS = (check_hidden_text, check_stuffing)

#: Which check set fits which kind of document, for the callers that know the
#: kind but nothing about the checks.
CHECKS_BY_KIND: dict[DocumentKind, tuple[Check, ...]] = {
    DocumentKind.RESUME: TEXT_CHECKS,
    DocumentKind.COVER_LETTER: LETTER_CHECKS,
}


def _word_count(pages: Sequence[PageFacts]) -> int:
    """Words in the document, however the pages were built.

    ``words`` is populated from geometry for a PDF and from the paragraphs for a
    DOCX; a plain-text page has none and is counted from its text.
    """
    counted = sum(len(page.words) for page in pages)
    return counted or sum(len(page.text.split()) for page in pages)


def _docx_pages(content: bytes, raw_text: str) -> list[PageFacts]:
    """Facts a DOCX can answer for.

    There is no geometry to read, but there is one thing a DOCX gets wrong more
    often than a PDF: layout built out of tables. Word makes that the path of
    least resistance, and many parsers read a table cell by cell, so a skills
    grid comes out as a shuffled bag of words. Counting table words here is what
    turns "we could not check the layout" into an answer.
    """
    document = docx.Document(io.BytesIO(content))
    in_tables = [
        word
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for word in cell.text.split()
    ]
    loose = [word for para in document.paragraphs for word in para.text.split()]
    words = [{"x0": 0.0, "top": 0.0, "text": word} for word in loose + in_tables]
    return [
        PageFacts(
            words=words,
            text=raw_text,
            width=0.0,
            columns=[],
            mixed_lines=[],
            line_count=len(raw_text.splitlines()),
            table_words=len(in_tables),
            chars=[],
        )
    ]


def _text_only_pages(text: str) -> list[PageFacts]:
    """Wrap plain text so the text-level checks can run on a DOCX or a TXT."""
    return [
        PageFacts(
            words=[],
            text=text,
            width=0.0,
            columns=[],
            mixed_lines=[],
            line_count=len(text.splitlines()),
            table_words=0,
            chars=[],
        )
    ]


# ── the two readings ──────────────────────────────────────────────────
#
# Everything above judges the file. This judges the *loss*: the structure the
# model read out of the resume, against the structure still recoverable from
# the text layer alone. The model saw the page and got it right; the text layer
# is all an employer's parser is given. The gap between them is the number the
# candidate actually came for — "an ATS will see 2 of your 4 jobs" — and it
# costs no extra model call, because the extraction has already been paid for.

#: Month names as resumes write them, by the prefix that identifies them.
MONTH_NAMES: dict[str, int] = {
    "янв": 1,
    "фев": 2,
    "мар": 3,
    "апр": 4,
    "мая": 5,
    "май": 5,
    "июн": 6,
    "июл": 7,
    "авг": 8,
    "сен": 9,
    "окт": 10,
    "ноя": 11,
    "дек": 12,
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}

#: Characters that differ between how a resume prints a name and how the model
#: reads it back. Stripped from both sides before comparing.
NOISE = re.compile(r"[«»\"'`()\[\]{},.;:!?—–\-\s]+")


def _flatten(value: str) -> str:
    """Casefold and drop punctuation, so two spellings of one name compare equal."""
    return NOISE.sub(" ", value).casefold().strip()


def year_month(token: str) -> tuple[int, int] | None:
    """Normalise one date as the text layer writes it to ``(year, month)``."""
    digits = re.findall(r"\d+", token)
    if len(digits) == 2:
        first, second = int(digits[0]), int(digits[1])
        if first > MONTHS_IN_YEAR:  # YYYY-MM
            return (first, second) if 1 <= second <= MONTHS_IN_YEAR else None
        return (second, first) if second > MONTHS_IN_YEAR else None
    if len(digits) == 1:
        head = _flatten(token)[:3]
        month = MONTH_NAMES.get(head)
        return (int(digits[0]), month) if month else None
    return None


def readable_lines(pages: Sequence[PageFacts]) -> list[str]:
    """Lines a sectioning parser would attribute to the right place.

    A line that mixes two columns is dropped, and so is one carrying a section
    heading somewhere other than its start: both are lines where the words are
    present but filed under the wrong job, which is worse than losing them —
    the candidate's freelance stint gets read as part of their skills list.
    """
    dropped = {line for page in pages for line in page.mixed_lines}
    clean: list[str] = []
    for page in pages:
        for line in page.text.splitlines():
            stripped = line.strip()
            if not stripped or stripped in dropped:
                continue
            heading = next(
                (m for p in SECTION_PATTERNS.values() if (m := p.search(stripped))), None
            )
            if heading is not None and heading.start() > 0:
                continue
            clean.append(stripped)
    return clean


def _recoverable(items: Sequence[str], lines: Sequence[str]) -> Recoverable:
    """How many of these survive into text a parser reads correctly."""
    haystack = " " + " ".join(_flatten(line) for line in lines) + " "
    named = [item for item in items if item and item.strip()]
    lost = [item for item in named if _flatten(item) not in haystack]
    return Recoverable(total=len(named), recovered=len(named) - len(lost), lost=lost[:10])


def _dates_recoverable(extraction: ProfileExtraction, lines: Sequence[str]) -> Recoverable:
    """Employment periods whose start date survives as a readable date.

    Compared as (year, month) rather than as strings: the model normalises every
    date to YYYY-MM and the resume prints 04.2022, «апрель 2022» or 2022-04, and
    all three mean the same thing to an ATS.
    """
    present = {
        parsed
        for line in lines
        for token in DATE_TOKEN.findall(line)
        if (parsed := year_month(token))
    }
    dated = [period for period in extraction.work_periods if period.start]
    lost = [
        f"{period.company}: {period.start}"
        for period in dated
        if (ym := year_month(str(period.start))) is None or ym not in present
    ]
    return Recoverable(total=len(dated), recovered=len(dated) - len(lost), lost=lost[:10])


def measure_coverage(extraction: ProfileExtraction, pages: Sequence[PageFacts]) -> ATSCoverage:
    """Compare what the model read with what a text-layer parser recovers."""
    lines = readable_lines(pages)
    return ATSCoverage(
        work_periods=_recoverable([p.company for p in extraction.work_periods], lines),
        dates=_dates_recoverable(extraction, lines),
        skills=_recoverable([skill.name for skill in extraction.skills], lines),
    )


def _coverage_finding(coverage: ATSCoverage) -> Finding | None:
    """Report the loss, in the terms the candidate cares about."""
    if coverage.is_complete:
        return None

    jobs, skills = coverage.work_periods, coverage.skills
    parts = [
        f"мест работы — {part.recovered} из {part.total}"
        if label == "jobs"
        else f"{label} — {part.recovered} из {part.total}"
        for label, part in (
            ("jobs", jobs),
            ("дат", coverage.dates),
            ("навыков", skills),
        )
        if not part.is_complete
    ]
    example = "; ".join(jobs.lost or coverage.dates.lost or skills.lost)
    return _finding(
        FindingCode.CONTENT_LOST,
        Severity.CRITICAL if jobs.recovered == 0 and jobs.total else Severity.WARNING,
        "Часть резюме не доходит до робота",
        "Из того, что видно человеку на странице, машинному чтению доступно: "
        + ", ".join(parts)
        + ". Остальное либо теряется, либо приклеивается к чужому разделу — "
        "робот отнесёт этот опыт не туда, где он есть.",
        "Убери двухколоночную вёрстку и таблицы: почти всегда расхождение "
        "берётся оттуда. После правки прогони проверку ещё раз — число должно "
        "стать полным.",
        example[:200] or None,
    )


def audit(
    content: bytes,
    *,
    source_format: str,
    raw_text: str = "",
    extraction: ProfileExtraction | None = None,
    keywords: ATSKeywords | None = None,
) -> ATSReport:
    """Judge how a machine will read this file.

    A PDF gets every check. A DOCX gets the text-level ones plus the table
    check, which is the layout mistake DOCX resumes actually make. Anything
    else gets the text-level ones and says so, rather than reporting a clean
    bill of health it did not earn.

    ``extraction`` is what the model read out of the same file. Passing it turns
    on the half of the report that matters most — how much of that survives into
    the text layer — and it is optional because the report is produced twice:
    once at upload, when no extraction exists yet and the structural findings
    are already worth showing, and again when parsing finishes.

    ``keywords`` is this document read against one vacancy's requirement list,
    computed by :mod:`app.resume.ats_keywords`. Optional for the same reason:
    the report shown right after upload is not about any particular vacancy.
    """
    checks: tuple[Check, ...]
    extra: list[Finding] = []
    if source_format == "pdf":
        pages = gather(content)
        checks = PDF_CHECKS
    else:
        docx_file = source_format == "docx"
        pages = _docx_pages(content, raw_text) if docx_file else _text_only_pages(raw_text)
        checks = DOCX_CHECKS if docx_file else TEXT_CHECKS
        extra = [
            _finding(
                FindingCode.FORMAT_NOT_PDF,
                Severity.INFO,
                f"Проверен как {source_format.upper()}, не как PDF",
                "Проверки вёрстки, которым нужна геометрия страницы — колонки, "
                "шрифты, глифы — относятся к PDF и здесь не выполнялись. "
                "Большинство работодателей ждёт PDF.",
                "Если отправляешь в PDF — прогони проверку ещё раз на самом PDF: "
                "вёрстка при экспорте меняется.",
            )
        ]

    return _assemble(
        pages,
        checks,
        extra=extra,
        extraction=extraction,
        keywords=keywords,
        kind=DocumentKind.RESUME,
        origin=DocumentOrigin.UPLOADED,
        source_format=source_format,
        page_count=len(pages) if source_format == "pdf" else 0,
    )


def audit_generated(
    text: str,
    *,
    kind: DocumentKind,
    keywords: ATSKeywords | None = None,
) -> ATSReport:
    """Audit a document this system produced, before a person sees it.

    The original audit answers a question about somebody else's file. This one
    turns the same machinery on our own output, which is the harder half of the
    idea: a generator is graded on the coverage number it produces, and the
    cheapest way to raise that number is exactly what
    :func:`check_stuffing` and :func:`check_hidden_text` refuse. Running them
    over every generated document is what keeps the boundary from being a
    paragraph in a prompt.

    The checks that apply are decided by ``kind`` rather than by the caller, so
    "audit what we wrote" cannot quietly become "audit it with the checks it
    happens to pass".

    Text only. Nothing here has a page yet — a letter is pasted into a form, and
    a generated resume is a document tree until something renders it — so the
    geometry checks have nothing to read and are honestly absent from
    ``checks_run`` rather than silently passing.
    """
    pages = _text_only_pages(text)
    return _assemble(
        pages,
        CHECKS_BY_KIND[kind],
        extra=[],
        extraction=None,
        keywords=keywords,
        kind=kind,
        origin=DocumentOrigin.GENERATED,
        source_format="text",
        page_count=0,
    )


def _assemble(
    pages: Sequence[PageFacts],
    checks: Sequence[Check],
    *,
    extra: Sequence[Finding],
    extraction: ProfileExtraction | None,
    keywords: ATSKeywords | None,
    kind: DocumentKind,
    origin: DocumentOrigin,
    source_format: str,
    page_count: int,
) -> ATSReport:
    """Run a check set over prepared pages and build the report.

    One place, so an uploaded resume and a document this system generated are
    scored by the same arithmetic. Two reports about the same defect that
    disagree about what it costs would make the number meaningless on both.
    """
    findings = [*extra]
    for check in checks:
        found = check(pages)
        if found is not None:
            findings.append(found)

    coverage: ATSCoverage | None = None
    if extraction is not None:
        coverage = measure_coverage(extraction, pages)
        for found in (_coverage_finding(coverage), check_gaps(extraction)):
            if found is not None:
                findings.append(found)

    if keywords is not None:
        unnamed = check_requirements(keywords)
        if unnamed is not None:
            findings.append(unnamed)

    # A file with no text layer cannot fail the checks that read the text; they
    # would all fire at once and bury the one finding that matters.
    if any(f.code is FindingCode.NO_TEXT_LAYER for f in findings):
        findings = [f for f in findings if f.code is FindingCode.NO_TEXT_LAYER]

    report = ATSReport(
        score=max(0, 100 - sum(f.penalty for f in findings)),
        findings=findings,
        checks_run=_codes_for(
            checks, any(f.code is FindingCode.FORMAT_NOT_PDF for f in extra), coverage, keywords
        ),
        sections_detected=detect_sections(pages),
        coverage=coverage,
        keywords=keywords,
        document_kind=kind,
        origin=origin,
        source_format=source_format,
        page_count=page_count,
        word_count=_word_count(pages),
    )
    logger.info(
        "resume.ats_audited",
        source_format=source_format,
        document_kind=kind.value,
        origin=origin.value,
        score=report.score,
        overall=report.overall.value,
        critical=len(report.critical),
        findings=len(report.findings),
        jobs_recovered=coverage.work_periods.recovered if coverage else None,
        jobs_total=coverage.work_periods.total if coverage else None,
        requirements=len(keywords.requirements) if keywords else None,
        requirements_present=len(keywords.present) if keywords else None,
    )
    return report


#: Which finding each check can raise, so a report can say what it looked at.
CHECK_CODES: dict[Check, FindingCode] = {
    check_text_layer: FindingCode.NO_TEXT_LAYER,
    check_columns: FindingCode.COLUMN_INTERLEAVING,
    check_unmapped_fonts: FindingCode.UNMAPPED_FONT,
    check_glyphs: FindingCode.BROKEN_GLYPHS,
    check_contacts: FindingCode.CONTACTS_NOT_TEXT,
    check_dates: FindingCode.DATES_NOT_EXTRACTABLE,
    check_tables: FindingCode.TEXT_IN_TABLES,
    check_sections: FindingCode.MISSING_SECTIONS,
    check_hidden_text: FindingCode.HIDDEN_TEXT,
    check_stuffing: FindingCode.KEYWORD_STUFFING,
    check_date_format: FindingCode.DATE_FORMAT_MIXED,
    check_length: FindingCode.LENGTH_OUT_OF_RANGE,
}


def _codes_for(
    checks: Sequence[Check],
    format_noted: bool,
    coverage: ATSCoverage | None,
    keywords: ATSKeywords | None,
) -> list[FindingCode]:
    """The codes a given run was actually able to look for.

    The two comparisons are listed only when they were made. A client reading
    this list is deciding what to render as "checked and clean", and a report
    that claims to have compared a document against a vacancy it never saw is
    worse than one that admits it did not.
    """
    codes = [CHECK_CODES[check] for check in checks]
    if format_noted:
        codes.append(FindingCode.FORMAT_NOT_PDF)
    if coverage is not None:
        codes.extend((FindingCode.CONTENT_LOST, FindingCode.EMPLOYMENT_GAP))
    if keywords is not None:
        codes.append(FindingCode.REQUIREMENTS_NOT_NAMED)
    return codes


def median_column_gap(pages: Sequence[PageFacts]) -> float:
    """Typical gutter width, for diagnostics. Zero when there is one column."""
    gaps = [
        page.columns[i + 1].left - page.columns[i].right
        for page in pages
        for i in range(len(page.columns) - 1)
    ]
    return statistics.median(gaps) if gaps else 0.0
